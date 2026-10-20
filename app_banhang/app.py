import os
import sqlite3
import calendar
from datetime import datetime
from flask import Blueprint, Flask, render_template, request, jsonify, send_file, redirect

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_NAME = os.path.join(BASE_DIR, "tiem_lap_do.db")
banhang_bp = Blueprint('banhang', __name__, template_folder='templates')

def init_db():
    with sqlite3.connect(DB_NAME, timeout=10) as conn:
        conn.execute("PRAGMA journal_mode = WAL;")
        conn.execute("PRAGMA synchronous = NORMAL;")
        c = conn.cursor()
        # Bảng 1: Đơn hàng
        c.execute("""CREATE TABLE IF NOT EXISTS orders (
                        id INTEGER PRIMARY KEY AUTOINCREMENT, 
                        customer_name TEXT,
                        phone TEXT, 
                        license_plate TEXT, 
                        install_date TEXT, 
                        total_price REAL)""")
        # Bảng 2: Chi tiết món đồ chơi theo từng đơn (Lưu đủ Đơn giá gốc, Giá sau giảm, Số lượng)
        c.execute("""CREATE TABLE IF NOT EXISTS order_items (
                        id INTEGER PRIMARY KEY AUTOINCREMENT, 
                        order_id INTEGER, 
                        item_name TEXT,
                        price REAL, 
                        warranty_months INTEGER,
                        orig_price REAL DEFAULT 0,
                        unit_price REAL DEFAULT 0,
                        qty INTEGER DEFAULT 1,
                        FOREIGN KEY(order_id) REFERENCES orders(id))""")
        c.execute("PRAGMA table_info(order_items)")
        oi_cols = [r[1] for r in c.fetchall()]
        if "orig_price" not in oi_cols: c.execute("ALTER TABLE order_items ADD COLUMN orig_price REAL DEFAULT 0")
        if "unit_price" not in oi_cols: c.execute("ALTER TABLE order_items ADD COLUMN unit_price REAL DEFAULT 0")
        if "qty" not in oi_cols: c.execute("ALTER TABLE order_items ADD COLUMN qty INTEGER DEFAULT 1")

        # Bảng 3: Kho hàng hóa chuẩn 100% theo Sheet HÀNG HÓA trong Excel
        c.execute("""CREATE TABLE IF NOT EXISTS products (
                        id INTEGER PRIMARY KEY AUTOINCREMENT,
                        item_type TEXT DEFAULT 'Hàng hóa',
                        category TEXT DEFAULT '',
                        code TEXT UNIQUE,
                        name TEXT,
                        price REAL DEFAULT 0,
                        cost_price REAL DEFAULT 0,
                        stock INTEGER DEFAULT 0,
                        warranty_months INTEGER DEFAULT 0,
                        warranty_raw TEXT DEFAULT '',
                        created_at TEXT DEFAULT '')""")
        # Nâng cấp các cột nếu bảng cũ chưa có
        c.execute("PRAGMA table_info(products)")
        cols = [r[1] for r in c.fetchall()]
        if "item_type" not in cols: c.execute("ALTER TABLE products ADD COLUMN item_type TEXT DEFAULT 'Hàng hóa'")
        if "category" not in cols: c.execute("ALTER TABLE products ADD COLUMN category TEXT DEFAULT ''")
        if "code" not in cols: c.execute("ALTER TABLE products ADD COLUMN code TEXT")
        if "cost_price" not in cols: c.execute("ALTER TABLE products ADD COLUMN cost_price REAL DEFAULT 0")
        if "stock" not in cols: c.execute("ALTER TABLE products ADD COLUMN stock INTEGER DEFAULT 0")
        if "warranty_raw" not in cols: c.execute("ALTER TABLE products ADD COLUMN warranty_raw TEXT DEFAULT ''")
        if "created_at" not in cols: c.execute("ALTER TABLE products ADD COLUMN created_at TEXT DEFAULT ''")

        # Tối ưu truy vấn tốc độ cao (Indexes)
        c.execute("CREATE INDEX IF NOT EXISTS idx_orders_plate ON orders(license_plate)")
        c.execute("CREATE INDEX IF NOT EXISTS idx_orders_phone ON orders(phone)")
        c.execute("CREATE INDEX IF NOT EXISTS idx_orders_date ON orders(install_date)")
        c.execute("CREATE INDEX IF NOT EXISTS idx_order_items_oid ON order_items(order_id)")
        c.execute("CREATE INDEX IF NOT EXISTS idx_products_name ON products(name)")
        c.execute("CREATE INDEX IF NOT EXISTS idx_products_code ON products(code)")
        c.execute("CREATE INDEX IF NOT EXISTS idx_products_stock ON products(stock)")
        conn.commit()

def add_months(sourcedate, months):
    month = sourcedate.month - 1 + months
    year = sourcedate.year + month // 12
    month = month % 12 + 1
    # Dùng thư viện calendar để tránh lỗi ngày 31 nhảy sang tháng thiếu
    day = min(sourcedate.day, calendar.monthrange(year, month)[1])
    return datetime(year, month, day)

# Khởi tạo database ngay khi chạy phần mềm
init_db()

@banhang_bp.after_request
def add_no_cache_headers(response):
    response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
    response.headers["Pragma"] = "no-cache"
    response.headers["Expires"] = "0"
    return response

@banhang_bp.route("/")
def index():
    return render_template("banhang_index.html")

# ==========================================
# API 1: LẤY BẢNG GIÁ, TỒN KHO & THỐNG KÊ CHUẨN EXCEL
# ==========================================
@banhang_bp.route("/api/products", methods=["GET"])
def get_products():
    q = request.args.get("q", "").strip()
    stock_filter = request.args.get("filter", "all").strip()
    sort_by = request.args.get("sort", "excel").strip()
    try:
        page = max(1, int(request.args.get("page", 1)))
    except Exception:
        page = 1
    per_page = 100

    with sqlite3.connect(DB_NAME) as conn:
        c = conn.cursor()
        
        # Thống kê tổng quan khớp 100% với Sheet HÀNG HÓA trong Excel
        try:
            c.execute("""SELECT 
                            COUNT(*),
                            COALESCE(SUM(stock), 0),
                            SUM(CASE WHEN stock > 0 THEN 1 ELSE 0 END),
                            SUM(CASE WHEN stock <= 0 AND item_type = 'Hàng hóa' THEN 1 ELSE 0 END),
                            SUM(CASE WHEN item_type = 'Dịch vụ' THEN 1 ELSE 0 END)
                         FROM products""")
            st = c.fetchone()
            stats = {
                "total": st[0] or 0,
                "total_stock": st[1] or 0,
                "in_stock": st[2] or 0,
                "out_of_stock": st[3] or 0,
                "service_count": st[4] or 0
            }
        except Exception:
            stats = {"total": 0, "total_stock": 0, "in_stock": 0, "out_of_stock": 0, "service_count": 0}

        where_clauses = []
        params = []
        
        if q:
            where_clauses.append("(name LIKE ? OR code LIKE ? OR category LIKE ?)")
            params.extend([f"%{q}%", f"%{q}%", f"%{q}%"])
            
        if stock_filter == "in":
            where_clauses.append("stock > 0")
        elif stock_filter == "out":
            where_clauses.append("stock <= 0 AND item_type = 'Hàng hóa'")
        elif stock_filter == "service":
            where_clauses.append("item_type = 'Dịch vụ'")

        where_sql = (" WHERE " + " AND ".join(where_clauses)) if where_clauses else ""

        if sort_by == "stock_desc":
            order_sql = " ORDER BY stock DESC, id ASC"
        elif sort_by == "stock_asc":
            order_sql = " ORDER BY stock ASC, id ASC"
        elif sort_by == "price_desc":
            order_sql = " ORDER BY price DESC, id ASC"
        else:
            # Mặc định giữ đúng 100% thứ tự dòng trong file Excel (từ dòng 2 trở xuống)
            order_sql = " ORDER BY id ASC"

        try:
            c.execute(f"SELECT COUNT(*) FROM products{where_sql}", params)
            total_matched = c.fetchone()[0] or 0
            total_pages = max(1, (total_matched + per_page - 1) // per_page)
            if page > total_pages:
                page = total_pages
            offset = (page - 1) * per_page

            c.execute(f"""SELECT id, item_type, category, code, name, price, cost_price, warranty_months, stock 
                          FROM products{where_sql}{order_sql} LIMIT ? OFFSET ?""", params + [per_page, offset])
            items = c.fetchall()
        except Exception:
            items = []
            total_matched = 0
            total_pages = 1
                
        result = [{
            "id": i[0],
            "type": i[1] or "Hàng hóa",
            "category": i[2] or "",
            "code": i[3] or "", 
            "name": i[4], 
            "price": i[5] or 0, 
            "cost": i[6] or 0, 
            "war": i[7] or 0, 
            "stock": i[8] if i[8] is not None else 0
        } for i in items]

    return jsonify({
        "products": result, 
        "stats": stats,
        "page": page,
        "total_pages": total_pages,
        "total_matched": total_matched
    })

# ==========================================
# API: LẤY TOÀN BỘ 2.764 MÓN ĐỂ TỰ ĐỘNG GỢI Ý Ở TAB NHẬP ĐƠN
# ==========================================
@banhang_bp.route("/api/products/all-suggestions", methods=["GET"])
def get_all_suggestions():
    with sqlite3.connect(DB_NAME) as conn:
        c = conn.cursor()
        c.execute("SELECT id, code, name, price, warranty_months, stock, item_type FROM products ORDER BY id ASC")
        rows = c.fetchall()
        items = [{
            "id": r[0],
            "code": r[1] or "",
            "name": r[2],
            "price": r[3] or 0,
            "war": r[4] or 0,
            "stock": r[5] if r[5] is not None else 0,
            "type": r[6] or "Hàng hóa"
        } for r in rows]
    return jsonify(items)

# ==========================================
# API: ĐỒNG BỘ LẠI TỒN KHO TRỰC TIẾP TỪ FILE EXCEL
# ==========================================
@banhang_bp.route("/api/products/sync-excel", methods=["POST"])
def sync_excel_api():
    try:
        from import_hang_hoa import sync_from_excel
        res = sync_from_excel()
        return jsonify(res)
    except Exception as e:
        return jsonify({"status": "error", "msg": f"Lỗi đồng bộ Excel: {str(e)}"}), 500

# ==========================================
# API: CẬP NHẬT TỒN KHO
# ==========================================
@banhang_bp.route("/api/products/update-stock", methods=["POST"])
def update_stock():
    data = request.json or {}
    pid = data.get("id")
    new_stock = data.get("stock")
    delta = data.get("delta")

    if not pid:
        return jsonify({"status": "error", "msg": "Thiếu mã sản phẩm"}), 400

    with sqlite3.connect(DB_NAME) as conn:
        c = conn.cursor()
        if new_stock is not None:
            c.execute("UPDATE products SET stock = ? WHERE id = ?", (int(new_stock), pid))
        elif delta is not None:
            c.execute("UPDATE products SET stock = stock + ? WHERE id = ?", (int(delta), pid))
        conn.commit()
        c.execute("SELECT stock, name, code FROM products WHERE id = ?", (pid,))
        row = c.fetchone()
        curr_stock = row[0] if row else 0
        pname = row[1] if row else ""
        pcode = row[2] if row else ""
    return jsonify({"status": "success", "msg": f"Đã cập nhật [{pcode}] {pname}: Tồn kho = {curr_stock}"})

# ==========================================
# API: THÊM MÓN MỚI VÀO KHO
# ==========================================
@banhang_bp.route("/api/products/add", methods=["POST"])
def add_product():
    data = request.json or {}
    name = data.get("name", "").strip()
    if not name:
        return jsonify({"status": "error", "msg": "Tên sản phẩm không được trống"}), 400
        
    code = data.get("code", "").strip().upper()
    category = data.get("category", "").strip()
    item_type = data.get("item_type", "Hàng hóa").strip()
    try: price = float(data.get("price", 0))
    except: price = 0.0
    try: cost = float(data.get("cost_price", 0))
    except: cost = 0.0
    try: war = int(data.get("warranty_months", 0))
    except: war = 0
    try: stock = int(data.get("stock", 0))
    except: stock = 0

    with sqlite3.connect(DB_NAME) as conn:
        c = conn.cursor()
        try:
            if not code:
                c.execute("SELECT COUNT(*) FROM products")
                cnt = (c.fetchone()[0] or 0) + 1
                code = f"SP_MOI_{cnt}"
            c.execute("""INSERT INTO products (item_type, category, code, name, price, cost_price, warranty_months, stock, created_at)
                         VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                         ON CONFLICT(code) DO UPDATE SET 
                            name=excluded.name,
                            category=excluded.category,
                            item_type=excluded.item_type,
                            price=excluded.price, 
                            cost_price=excluded.cost_price, 
                            warranty_months=excluded.warranty_months, 
                            stock=excluded.stock""", 
                      (item_type, category, code, name, price, cost, war, stock, datetime.now().strftime("%Y-%m-%d %H:%M:%S")))
            conn.commit()
            return jsonify({"status": "success", "msg": f"Đã lưu [{code}] {name} (Tồn: {stock})"})
        except Exception as e:
            return jsonify({"status": "error", "msg": str(e)}), 500

# ==========================================
# API: SỬA GIÁ BÁN, TỒN KHO & THÔNG TIN HÀNG HÓA
# ==========================================
@banhang_bp.route("/api/products/edit", methods=["POST"])
def edit_product():
    data = request.json or {}
    pid = data.get("id")
    if not pid:
        return jsonify({"status": "error", "msg": "Thiếu ID sản phẩm cần sửa"}), 400

    with sqlite3.connect(DB_NAME) as conn:
        c = conn.cursor()
        c.execute("SELECT id, code, name, price, cost_price, stock, warranty_months, category, item_type FROM products WHERE id = ?", (pid,))
        row = c.fetchone()
        if not row:
            return jsonify({"status": "error", "msg": "Không tìm thấy sản phẩm trong kho"}), 404

        new_code = data.get("code", row[1] or "").strip().upper() or row[1]
        new_name = data.get("name", row[2] or "").strip() or row[2]
        try: new_price = float(data["price"]) if "price" in data and data["price"] is not None else float(row[3] or 0)
        except: new_price = float(row[3] or 0)
        try: new_cost = float(data["cost_price"]) if "cost_price" in data and data["cost_price"] is not None else float(row[4] or 0)
        except: new_cost = float(row[4] or 0)
        try: new_stock = int(data["stock"]) if "stock" in data and data["stock"] is not None else int(row[5] or 0)
        except: new_stock = int(row[5] or 0)
        try: new_war = int(data["warranty_months"]) if "warranty_months" in data and data["warranty_months"] is not None else int(row[6] or 0)
        except: new_war = int(row[6] or 0)
        new_cat = data.get("category", row[7] or "").strip()
        new_type = data.get("item_type", row[8] or "Hàng hóa").strip()

        try:
            c.execute("""UPDATE products 
                         SET code = ?, name = ?, price = ?, cost_price = ?, stock = ?, warranty_months = ?, category = ?, item_type = ?
                         WHERE id = ?""",
                      (new_code, new_name, new_price, new_cost, new_stock, new_war, new_cat, new_type, pid))
            conn.commit()
            return jsonify({
                "status": "success",
                "msg": f"Đã cập nhật [{new_code}] {new_name}: Giá bán {new_price:,.0f}đ | Tồn kho: {new_stock}"
            })
        except Exception as e:
            return jsonify({"status": "error", "msg": f"Lỗi khi cập nhật: {str(e)}"}), 500

@banhang_bp.route("/api/products/delete", methods=["POST"])
def delete_product():
    data = request.json or {}
    pid = data.get("id")
    if not pid:
        return jsonify({"status": "error", "msg": "Thiếu ID sản phẩm"}), 400
    with sqlite3.connect(DB_NAME) as conn:
        c = conn.cursor()
        c.execute("DELETE FROM products WHERE id = ?", (pid,))
        conn.commit()
    return jsonify({"status": "success", "msg": "Đã xóa mặt hàng khỏi kho!"})

# ==========================================
# API 2: LƯU ĐƠN HÀNG MỚI (TỰ ĐỘNG TRỪ TỒN KHO)
# ==========================================
@banhang_bp.route("/api/order", methods=["POST"])
def save_order():
    data = request.json
    plate = data.get("plate", "").strip().upper()
    name = data.get("name", "").strip()
    phone = data.get("phone", "").strip()
    date_str = data.get("date", "").strip()
    items = data.get("items", [])

    if not plate or not items:
        return jsonify({"status": "error", "msg": "Thiếu biển số hoặc món hàng"}), 400

    try:
        datetime.strptime(date_str, "%Y-%m-%d")
    except ValueError:
        return jsonify({"status": "error", "msg": "Ngày tháng sai định dạng YYYY-MM-DD"}), 400

    total = sum(float(i["price"]) for i in items)

    with sqlite3.connect(DB_NAME) as conn:
        c = conn.cursor()
        c.execute("""INSERT INTO orders (customer_name, phone, license_plate, install_date, total_price) 
                     VALUES (?, ?, ?, ?, ?)""", (name, phone, plate, date_str, total))
        order_id = c.lastrowid
        
        for it in items:
            price_val = float(it["price"])
            qty = max(1, int(it.get("qty", 1)))
            unit_price = float(it.get("unit_price", price_val / qty if qty else price_val))
            orig_price = float(it.get("orig_unit_price", unit_price))

            c.execute("""INSERT INTO order_items (order_id, item_name, price, warranty_months, orig_price, unit_price, qty) 
                         VALUES (?, ?, ?, ?, ?, ?, ?)""", 
                      (order_id, it["name"], price_val, int(it["warranty"]), orig_price, unit_price, qty))
            
            # TỰ ĐỘNG TRỪ TỒN KHO: Chỉ trừ nếu là Hàng hóa (đơn giá >= 0)
            if price_val >= 0:
                prod_id = it.get("product_id")
                if prod_id:
                    c.execute("UPDATE products SET stock = stock - ? WHERE id = ? AND item_type = 'Hàng hóa'", (qty, prod_id))
                else:
                    clean_name = (it.get("raw_name") or it["name"].split(" [Gốc")[0].split(" (Giảm")[0]).strip()
                    c.execute("""UPDATE products SET stock = stock - ? 
                                 WHERE id = (SELECT id FROM products WHERE name = ? AND item_type = 'Hàng hóa' ORDER BY id ASC LIMIT 1)""", (qty, clean_name))
        conn.commit()
    return jsonify({"status": "success", "msg": f"Đã lưu đơn hàng (Ngày ra đơn: {date_str}) & cập nhật tồn kho!"})

# ==========================================
# API: SỬA THÔNG TIN KHÁCH HÀNG & ĐƠN HÀNG
# ==========================================
@banhang_bp.route("/api/order/update", methods=["POST"])
def update_order():
    data = request.json or {}
    oid = data.get("id")
    if not oid:
        return jsonify({"status": "error", "msg": "Thiếu mã đơn hàng cần sửa"}), 400

    plate = data.get("plate", "").strip().upper()
    name = data.get("name", "").strip()
    phone = data.get("phone", "").strip()
    date_str = data.get("date", "").strip()
    items = data.get("items")

    if not plate:
        return jsonify({"status": "error", "msg": "Biển số xe không được để trống"}), 400

    with sqlite3.connect(DB_NAME) as conn:
        c = conn.cursor()
        if items is not None and isinstance(items, list) and len(items) > 0:
            total = sum(float(i.get("price", 0)) for i in items)
            c.execute("""UPDATE orders 
                         SET customer_name = ?, phone = ?, license_plate = ?, install_date = ?, total_price = ?
                         WHERE id = ?""", (name, phone, plate, date_str, total, oid))
            c.execute("DELETE FROM order_items WHERE order_id = ?", (oid,))
            for it in items:
                p_val = float(it.get("price", 0))
                q_val = max(1, int(it.get("qty", 1)))
                u_val = float(it.get("unit_price", p_val / q_val if q_val else p_val))
                o_val = float(it.get("orig_price", u_val))
                w_val = int(it.get("war", 0))
                c.execute("""INSERT INTO order_items (order_id, item_name, price, warranty_months, orig_price, unit_price, qty)
                             VALUES (?, ?, ?, ?, ?, ?, ?)""",
                          (oid, it.get("name", "").strip(), p_val, w_val, o_val, u_val, q_val))
        else:
            try:
                total = float(data.get("total"))
                c.execute("""UPDATE orders 
                             SET customer_name = ?, phone = ?, license_plate = ?, install_date = ?, total_price = ?
                             WHERE id = ?""", (name, phone, plate, date_str, total, oid))
            except Exception:
                c.execute("""UPDATE orders 
                             SET customer_name = ?, phone = ?, license_plate = ?, install_date = ?
                             WHERE id = ?""", (name, phone, plate, date_str, oid))
        conn.commit()
    return jsonify({"status": "success", "msg": f"Đã cập nhật thông tin khách hàng & đơn #{oid}!"})

@banhang_bp.route("/api/order/delete", methods=["POST"])
def delete_order():
    data = request.json or {}
    oid = data.get("id")
    if not oid:
        return jsonify({"status": "error", "msg": "Thiếu mã đơn hàng"}), 400
    with sqlite3.connect(DB_NAME) as conn:
        c = conn.cursor()
        c.execute("DELETE FROM order_items WHERE order_id = ?", (oid,))
        c.execute("DELETE FROM orders WHERE id = ?", (oid,))
        conn.commit()
    return jsonify({"status": "success", "msg": f"Đã xóa đơn hàng #{oid}!"})

# ==========================================
# API 3: TRA CỨU ĐƠN HÀNG & DANH SÁCH KHÁCH HÀNG (ĐÃ TỐI ƯU SIÊU TỐC)
# ==========================================
@banhang_bp.route("/api/search", methods=["GET"])
def search_orders():
    q = request.args.get("q", "").strip()
    date_filter = request.args.get("date", "").strip()
    with sqlite3.connect(DB_NAME) as conn:
        c = conn.cursor()
        
        where_clauses = []
        params = []
        if q:
            where_clauses.append("(license_plate LIKE ? OR phone LIKE ? OR customer_name LIKE ? OR install_date LIKE ?)")
            params.extend([f"%{q}%", f"%{q}%", f"%{q}%", f"%{q}%"])
        if date_filter:
            where_clauses.append("install_date = ?")
            params.append(date_filter)

        where_sql = (" WHERE " + " AND ".join(where_clauses)) if where_clauses else ""

        # 1. CHỐNG TREO MÁY: Giới hạn chỉ lấy tối đa 50 đơn hàng khớp nhất
        c.execute(f"""SELECT id, customer_name, phone, license_plate, install_date, total_price 
                      FROM orders{where_sql}
                      ORDER BY install_date DESC, id DESC LIMIT 50""", params)
        
        orders = c.fetchall()
        if not orders:
            return jsonify([])

        # 2. CHỐNG N+1 QUERY: Gom tất cả ID đơn hàng lại bằng 1 câu lệnh IN(...)
        order_ids = [str(o[0]) for o in orders]
        placeholders = ",".join("?" * len(order_ids))
        
        c.execute(f"""SELECT order_id, item_name, price, warranty_months, orig_price, unit_price, qty, id 
                      FROM order_items WHERE order_id IN ({placeholders})""", order_ids)
        all_items = c.fetchall()

        # Dùng Dictionary để nhóm đồ chơi vào đúng ID đơn hàng cực nhanh
        items_dict = {}
        for it in all_items:
            oid = it[0]
            if oid not in items_dict:
                items_dict[oid] = []
            items_dict[oid].append(it)

        result = []
        today = datetime.now()
        
        for o in orders:
            oid = o[0]
            order_items = items_dict.get(oid, [])
            item_list = []
            raw_date = o[4] or ""
            
            try: 
                install_date = datetime.strptime(raw_date, "%Y-%m-%d")
                date_vn = install_date.strftime("%d/%m/%Y")
            except: 
                install_date = today
                date_vn = raw_date
                
            for it in order_items:
                war_m = it[3] or 0
                if war_m > 0:
                    exp_date = add_months(install_date, war_m)
                    status = "✅ Còn BH" if exp_date >= today else "❌ Hết BH"
                    exp_str = exp_date.strftime("%d/%m/%Y")
                else:
                    exp_str = "Không BH"
                    status = "-"
                price_val = it[2] or 0
                orig_val = it[4] if (it[4] and it[4] > 0) else price_val
                unit_val = it[5] if (it[5] and it[5] > 0) else price_val
                qty_val = it[6] if (it[6] and it[6] > 0) else 1
                item_list.append({
                    "item_id": it[7],
                    "name": it[1], 
                    "price": price_val, 
                    "orig_price": orig_val,
                    "unit_price": unit_val,
                    "qty": qty_val,
                    "war": war_m, 
                    "exp": exp_str, 
                    "status": status
                })
            
            result.append({
                "id": o[0], 
                "name": o[1], 
                "phone": o[2], 
                "plate": o[3], 
                "date": raw_date, 
                "date_vn": date_vn,
                "total": o[5] or 0, 
                "items": item_list
            })
            
    return jsonify(result)

# ==========================================
# API 4: HỆ THỐNG EXCEL & BACKUP
# ==========================================
@banhang_bp.route("/api/export", methods=["GET"])
def export_excel():
    from openpyxl.workbook import Workbook
    dest = os.path.join(BASE_DIR, "danh_sach_don.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Danh Sách Khách Hàng"
    ws.append(["Mã Đơn", "Ngày Ra Đơn Hàng", "Tên Khách Hàng", "Số Điện Thoại", "Biển Số Xe", "Tổng Tiền Thực Thu"])
    
    with sqlite3.connect(DB_NAME, timeout=10) as conn:
        c = conn.cursor()
        c.execute("SELECT id, install_date, customer_name, phone, license_plate, total_price FROM orders ORDER BY install_date DESC, id DESC")
        for r in c.fetchall(): ws.append(list(r))
    
    wb.save(dest)
    return send_file(dest, as_attachment=True)

@banhang_bp.route("/api/backup", methods=["GET"])
def backup_db():
    return send_file(DB_NAME, as_attachment=True)

@banhang_bp.route("/api/import", methods=["POST"])
def import_excel():
    if 'file' not in request.files:
        return jsonify({"status": "error", "msg": "Vui lòng chọn file Excel hoặc CSV!"}), 400
    file = request.files['file']
    filename = (file.filename or '').lower()
    
    # 1. HỖ TRỢ FILE XUẤT TỪ MEPOS (.csv)
    if filename.endswith(".csv"):
        try:
            content = file.read().decode("utf-8-sig", errors="ignore")
            from .import_mepos import import_mepos_invoices
            res = import_mepos_invoices(content)
            return jsonify(res)
        except Exception as e:
            return jsonify({"status": "error", "msg": f"Lỗi đọc file CSV: {str(e)}"}), 500

    # 2. HỖ TRỢ FILE EXCEL (.xlsx / .xls)
    try:
        import pandas as pd
        xls_df = pd.read_excel(file)
        cols = [str(c).strip() for c in xls_df.columns]
        
        # Nếu là file xuất hóa đơn của MePOS dạng Excel
        if any("Mã hóa đơn" in c for c in cols) and any("Tên khách hàng" in c for c in cols):
            csv_buf = xls_df.to_csv(index=False)
            from .import_mepos import import_mepos_invoices
            res = import_mepos_invoices(csv_buf)
            return jsonify(res)

        # Định dạng Excel truyền thống
        orders_dict = {}
        for _, row in xls_df.iterrows():
            r = list(row)
            if len(r) < 5 or pd.isna(r[0]) or pd.isna(r[4]):
                continue
            plate = str(r[0]).strip().upper()
            name = str(r[1] or "").strip()
            phone_raw = str(r[2] or "").strip()
            if phone_raw.endswith('.0'): phone_raw = phone_raw[:-2]
            if phone_raw and not phone_raw.startswith('0') and len(phone_raw) in (9, 10) and phone_raw.isdigit():
                phone_raw = '0' + phone_raw
            phone = phone_raw
            date_str = str(r[3] or datetime.now().strftime("%Y-%m-%d")).strip()
            item_name = str(r[4]).strip()
            price = float(r[5]) if len(r) > 5 and pd.notna(r[5]) else 0.0
            war = int(r[6]) if len(r) > 6 and pd.notna(r[6]) else 0
            
            key = (plate, date_str)
            if key not in orders_dict: 
                orders_dict[key] = {"name": name, "phone": phone, "items": []}
            orders_dict[key]["items"].append((item_name, price, war))

        with sqlite3.connect(DB_NAME) as conn:
            c = conn.cursor()
            inserted = 0
            for (plate, date_str), data in orders_dict.items():
                total = sum(i[1] for i in data["items"])
                c.execute("""INSERT INTO orders (customer_name, phone, license_plate, install_date, total_price) 
                             VALUES (?, ?, ?, ?, ?)""", 
                          (data["name"], data["phone"], plate, date_str, total))
                oid = c.lastrowid
                inserted += 1
                for it in data["items"]:
                    c.execute("""INSERT INTO order_items (order_id, item_name, price, warranty_months, orig_price, unit_price, qty) 
                                 VALUES (?, ?, ?, ?, ?, ?, ?)""", (oid, it[0], it[1], it[2], it[1], it[1], 1))
            conn.commit()
        return jsonify({"status": "success", "msg": f"Đã nhập thành công {inserted} đơn hàng từ file Excel!"})
    except Exception as e:
        return jsonify({"status": "error", "msg": f"Lỗi xử lý file Excel: {str(e)}"}), 500

# ==========================================
# API 5: TÊN MIỀN TRUY CẬP MẠNG NỘI BỘ (KHÔNG CẦN NHỚ IP)
# ==========================================
import socket
import threading

PORT_80_ACTIVE = False

def get_lan_ip():
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.connect(("8.8.8.8", 80))
        ip = s.getsockname()[0]
        s.close()
        return ip
    except Exception:
        return "127.0.0.1"

@banhang_bp.route("/api/network-info", methods=["GET"])
def network_info():
    ip = get_lan_ip()
    hostname = socket.gethostname().lower()
    port_suffix = "" if PORT_80_ACTIVE else ":5000"
    dashed_ip = ip.replace(".", "-")
    return jsonify({
        "lan_ip": ip,
        "port_80": PORT_80_ACTIVE,
        "primary_domain": f"http://tiemlapdo.local{port_suffix}",
        "short_domain": f"http://banhang.local{port_suffix}",
        "pc_domain": f"http://{hostname}.local{port_suffix}",
        "universal_domain": f"http://tiemlapdo.{dashed_ip}.sslip.io{port_suffix}",
        "ip_url": f"http://{ip}{port_suffix}",
        "port5000_domain": "http://tiemlapdo.local:5000"
    })

def start_mdns_and_port80():
    global PORT_80_ACTIVE
    ip = get_lan_ip()
    
    # 1. Mở thêm cổng 80 (HTTP mặc định) để gõ thẳng tên miền không cần :5000
    try:
        from werkzeug.serving import make_server
        srv80 = make_server("0.0.0.0", 80, app, threaded=True)
        t80 = threading.Thread(target=srv80.serve_forever, daemon=True)
        t80.start()
        PORT_80_ACTIVE = True
        print("🌐 Đã mở cổng 80: Không cần gõ :5000 ở cuối tên miền!")
    except Exception as e:
        print(f"ℹ️ Cổng 80 đang bận ({e}), dùng cổng 5000.")

    # 2. Phát sóng tên miền dễ nhớ 'tiemlapdo.local' & 'banhang.local' qua Wi-Fi (mDNS/Bonjour)
    try:
        from zeroconf import Zeroconf, ServiceInfo
        zc = Zeroconf()
        target_port = 80 if PORT_80_ACTIVE else 5000
        addr_bytes = [socket.inet_aton(ip)]
        
        info1 = ServiceInfo(
            "_http._tcp.local.",
            "TiemLapDo._http._tcp.local.",
            addresses=addr_bytes,
            port=target_port,
            server="tiemlapdo.local."
        )
        info2 = ServiceInfo(
            "_http._tcp.local.",
            "BanHang._http._tcp.local.",
            addresses=addr_bytes,
            port=target_port,
            server="banhang.local."
        )
        zc.register_service(info1)
        zc.register_service(info2)
        suffix = "" if PORT_80_ACTIVE else ":5000"
        print("==========================================================")
        print("🚀 TRUY CẬP TỪ ĐIỆN THOẠI / MÁY KHÁC CÙNG WI-FI:")
        print(f"   👉 Gõ tên miền 1: http://tiemlapdo.local{suffix}")
        print(f"   👉 Gõ tên miền 2: http://banhang.local{suffix}")
        print(f"   👉 Tên máy tính : http://{socket.gethostname().lower()}.local{suffix}")
        print("==========================================================")
    except Exception as e:
        print(f"Lưu ý mDNS: {e}")

app = Flask(__name__)
app.register_blueprint(banhang_bp, url_prefix='/banhang')

@app.route('/')
def root_redirect():
    return redirect('/banhang/')

if __name__ == "__main__":
    start_mdns_and_port80()
    app.run(host='0.0.0.0', port=5000, debug=False, threaded=True)